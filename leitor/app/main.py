"""Leitor de documentos: arquivo -> JSON no contrato contracts/leitor/v1.

Sem estado, sem banco e sem regra de negócio. Só devolve o conteúdo bruto com posição.
Quem interpreta o layout (fluxo de caixa, balancete, PO...) é o backend Java.
"""
import hashlib
import io
from datetime import date, datetime

import pdfplumber
from docx import Document
from fastapi import FastAPI, HTTPException, UploadFile
from openpyxl import load_workbook

VERSAO_LEITOR = "leitor-py 0.1.0"
app = FastAPI(title="Leitor de documentos", version="0.1.0")


@app.get("/saude")
def saude():
    return {"status": "ok", "leitor": VERSAO_LEITOR}


@app.post("/v1/ler")
async def ler(arquivo: UploadFile):
    conteudo = await arquivo.read()
    nome = arquivo.filename or "arquivo"
    tipo = detectar_tipo(conteudo, nome)
    resultado = {
        "versao_contrato": "1",
        "leitor": VERSAO_LEITOR,
        "arquivo": {
            "nome": nome,
            "sha256": hashlib.sha256(conteudo).hexdigest(),
            "tamanho_bytes": len(conteudo),
        },
        "tipo": tipo,
        "paginas": [],
        "planilhas": [],
        "paragrafos": [],
    }
    try:
        if tipo == "pdf":
            resultado["paginas"] = ler_pdf(conteudo)
        elif tipo == "xlsx":
            resultado["planilhas"] = ler_xlsx(conteudo)
        else:
            resultado["paragrafos"] = ler_docx(conteudo)
    except Exception as erro:  # arquivo corrompido ou protegido
        raise HTTPException(status_code=422, detail=f"Não foi possível ler o arquivo: {erro}")
    return resultado


def detectar_tipo(conteudo: bytes, nome: str) -> str:
    if conteudo.startswith(b"%PDF"):
        return "pdf"
    if conteudo.startswith(b"PK"):  # xlsx e docx são zip
        if b"word/" in conteudo[:4096] or nome.lower().endswith(".docx"):
            return "docx"
        return "xlsx"
    raise HTTPException(status_code=415, detail="Formato não suportado. Envie PDF, Excel (.xlsx) ou Word (.docx).")


def ler_pdf(conteudo: bytes) -> list:
    paginas = []
    with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
        for numero, pagina in enumerate(pdf.pages, start=1):
            # Alguns relatórios desenham cada caractere duas vezes (uma "sombra" branca deslocada, ou negrito
            # simulado). Sem isso, "21/09/2026" vira "2211//0099//22002266". O tamanho da fonte das duas cópias
            # difere na 15ª casa decimal, por isso a comparação usa só a fonte, não o tamanho.
            palavras = pagina.dedupe_chars(tolerance=1, extra_attrs=("fontname",)).extract_words(
                keep_blank_chars=False, y_tolerance=2)
            paginas.append({
                "numero": numero,
                "largura": round(float(pagina.width), 2),
                "altura": round(float(pagina.height), 2),
                "metodo": "texto" if palavras else "sem_texto",
                "palavras": [
                    {
                        "texto": p["text"],
                        "x0": round(float(p["x0"]), 2),
                        "x1": round(float(p["x1"]), 2),
                        "topo": round(float(p["top"]), 2),
                        "base": round(float(p["bottom"]), 2),
                    }
                    for p in palavras
                ],
            })
    return paginas


def ler_xlsx(conteudo: bytes) -> list:
    livro = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    planilhas = []
    for aba in livro.worksheets:
        celulas = []
        for linha in aba.iter_rows():
            for celula in linha:
                if celula.value is None or celula.value == "":
                    continue
                celulas.append({
                    "linha": celula.row,
                    "coluna": celula.column,
                    "valor": valor_texto(celula.value),
                    "tipo": tipo_valor(celula.value),
                })
        planilhas.append({"nome": aba.title, "celulas": celulas})
    return planilhas


def valor_texto(valor) -> str:
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    return str(valor)


def tipo_valor(valor) -> str:
    if isinstance(valor, bool):
        return "booleano"
    if isinstance(valor, (int, float)):
        return "numero"
    if isinstance(valor, (datetime, date)):
        return "data"
    return "texto"


def ler_docx(conteudo: bytes) -> list:
    documento = Document(io.BytesIO(conteudo))
    paragrafos = []
    ordem = 0
    for p in documento.paragraphs:
        if p.text.strip():
            ordem += 1
            paragrafos.append({"ordem": ordem, "texto": p.text})
    for indice, tabela in enumerate(documento.tables, start=1):
        for linha in tabela.rows:
            texto = "\t".join(c.text.strip() for c in linha.cells)
            if texto.strip():
                ordem += 1
                paragrafos.append({"ordem": ordem, "texto": texto, "tabela": indice})
    return paragrafos
