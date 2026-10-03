"""Garante que a saída do leitor respeita o contrato JSON Schema v1."""
import io
import json
from pathlib import Path

import jsonschema
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app

RAIZ = Path(__file__).resolve().parents[2]
ESQUEMA = json.loads((RAIZ / "contracts/leitor/v1/documento-lido.schema.json").read_text())
cliente = TestClient(app)


def enviar(nome: str, conteudo: bytes) -> dict:
    resposta = cliente.post("/v1/ler", files={"arquivo": (nome, conteudo)})
    assert resposta.status_code == 200, resposta.text
    corpo = resposta.json()
    jsonschema.validate(corpo, ESQUEMA)
    return corpo


def test_xlsx():
    livro = Workbook()
    aba = livro.active
    aba["A1"] = "Rubrica"
    aba["B1"] = 1234.56
    buffer = io.BytesIO()
    livro.save(buffer)
    corpo = enviar("po.xlsx", buffer.getvalue())
    assert corpo["tipo"] == "xlsx"
    assert {"linha": 1, "coluna": 2, "valor": "1234.56", "tipo": "numero"} in corpo["planilhas"][0]["celulas"]


def test_docx():
    documento = Document()
    documento.add_paragraph("Ata da assembleia")
    buffer = io.BytesIO()
    documento.save(buffer)
    corpo = enviar("ata.docx", buffer.getvalue())
    assert corpo["tipo"] == "docx"
    assert corpo["paragrafos"][0]["texto"] == "Ata da assembleia"


def test_pdf_de_exemplo_quando_disponivel():
    exemplo = RAIZ / "data/golden/privado/fluxo-caixa-2026-09.pdf"
    if not exemplo.exists():
        import pytest
        pytest.skip("PDF de exemplo não está no repositório (dado real)")
    corpo = enviar(exemplo.name, exemplo.read_bytes())
    assert corpo["tipo"] == "pdf"
    assert len(corpo["paginas"]) == 22


def test_formato_nao_suportado():
    resposta = cliente.post("/v1/ler", files={"arquivo": ("x.txt", b"ola")})
    assert resposta.status_code == 415
