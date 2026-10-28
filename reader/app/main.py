"""Document reader: file -> JSON in the contracts/reader/v2 contract.

No state, no database and no business rules. It only returns the raw content with positions.
The layout interpretation (cash flow, trial balance, budget...) is done by the rag service in Java.
"""
import hashlib
import io
import logging
import time
import zipfile
from datetime import date, datetime

import pdfplumber
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from fastapi import FastAPI, HTTPException, Request, UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pdfminer.pdfdocument import PDFEncryptionError, PDFPasswordIncorrect
from pdfminer.pdfparser import PDFSyntaxError
from pdfplumber.utils.exceptions import MalformedPDFException, PdfminerException

from app import logs

READER_VERSION = "reader-py 0.2.0"
app = FastAPI(title="Leitor de documentos", version="0.2.0")
logs.configure()
log = logging.getLogger("reader")

# A file that cannot be read for a reason of its own (corrupted, protected, not really a PDF/zip): expected, WARN.
# Anything else is a bug of the reader: ERROR with the stack (ADR 0008, B2). The answer is 422 in both cases.
FILE_PROBLEMS = (PdfminerException, MalformedPDFException, PDFSyntaxError, PDFPasswordIncorrect, PDFEncryptionError,
                 zipfile.BadZipFile, InvalidFileException, PackageNotFoundError)


@app.middleware("http")
async def request_log(request: Request, call_next):
    """Request id of the X-Request-Id header (or a new one) in every line and in the response; one line per request
    instead of uvicorn's access log, which has no id. Only method, route, status and time: never the file name."""
    token = logs.request_id.set(logs.accept_or_create(request.headers.get(logs.HEADER)))
    start = time.perf_counter()
    try:
        response = await call_next(request)
        response.headers[logs.HEADER] = logs.request_id.get()
        if request.url.path != "/health":
            log.info("%s %s → %d em %d ms", request.method, request.url.path, response.status_code,
                     (time.perf_counter() - start) * 1000)
        return response
    finally:
        logs.request_id.reset(token)


@app.get("/health")
def health():
    return {"status": "ok", "reader": READER_VERSION}


@app.post("/v2/read")
async def read(file: UploadFile):
    content = await file.read()
    name = file.filename or "arquivo"
    file_type = detect_type(content, name)
    result = {
        "contractVersion": "2",
        "reader": READER_VERSION,
        "file": {
            "name": name,
            "sha256": hashlib.sha256(content).hexdigest(),
            "sizeBytes": len(content),
        },
        "type": file_type,
        "pages": [],
        "sheets": [],
        "paragraphs": [],
    }
    try:
        if file_type == "pdf":
            result["pages"] = read_pdf(content)
        elif file_type == "xlsx":
            result["sheets"] = read_xlsx(content)
        else:
            result["paragraphs"] = read_docx(content)
    except FILE_PROBLEMS as error:  # corrupted or protected file
        log.warning("Arquivo %s (sha256 %s) não pôde ser lido: %s", file_type, result["file"]["sha256"][:12],
                    type(error).__name__)
        raise HTTPException(status_code=422, detail=f"Não foi possível ler o arquivo: {error}")
    except Exception as error:  # a bug of the reader: same answer, but with the stack in the log
        log.error("Erro inesperado ao ler arquivo %s (sha256 %s)", file_type, result["file"]["sha256"][:12],
                  exc_info=True)
        raise HTTPException(status_code=422, detail=f"Não foi possível ler o arquivo: {error}")
    log.info("Arquivo %s lido (sha256 %s): %d página(s), %d planilha(s), %d parágrafo(s)", file_type,
             result["file"]["sha256"][:12], len(result["pages"]), len(result["sheets"]), len(result["paragraphs"]))
    return result


def detect_type(content: bytes, name: str) -> str:
    if content.startswith(b"%PDF"):
        return "pdf"
    if content.startswith(b"PK"):  # xlsx and docx are zip files
        if b"word/" in content[:4096] or name.lower().endswith(".docx"):
            return "docx"
        return "xlsx"
    log.warning("Formato não suportado (%d bytes)", len(content))
    raise HTTPException(status_code=415, detail="Formato não suportado. Envie PDF, Excel (.xlsx) ou Word (.docx).")


def read_pdf(content: bytes) -> list:
    pages = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            # Some reports draw each character twice (a shifted white "shadow", or simulated bold).
            # Without this, "21/09/2026" becomes "2211//0099//22002266". The font size of the two copies
            # differs in the 15th decimal place, so the comparison uses only the font, not the size.
            words = page.dedupe_chars(tolerance=1, extra_attrs=("fontname",)).extract_words(
                keep_blank_chars=False, y_tolerance=2)
            pages.append({
                "number": number,
                "width": round(float(page.width), 2),
                "height": round(float(page.height), 2),
                "method": "text" if words else "no_text",
                "words": [
                    {
                        "text": w["text"],
                        "x0": round(float(w["x0"]), 2),
                        "x1": round(float(w["x1"]), 2),
                        "top": round(float(w["top"]), 2),
                        "bottom": round(float(w["bottom"]), 2),
                    }
                    for w in words
                ],
            })
    return pages


def read_xlsx(content: bytes) -> list:
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheets = []
    for worksheet in workbook.worksheets:
        cells = []
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None or cell.value == "":
                    continue
                cells.append({
                    "row": cell.row,
                    "column": cell.column,
                    "value": value_text(cell.value),
                    "type": value_type(cell.value),
                })
        sheets.append({"name": worksheet.title, "cells": cells})
    return sheets


def value_text(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def value_type(value) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (datetime, date)):
        return "date"
    return "text"


def read_docx(content: bytes) -> list:
    document = Document(io.BytesIO(content))
    paragraphs = []
    sequence = 0
    for p in document.paragraphs:
        if p.text.strip():
            sequence += 1
            paragraphs.append({"sequence": sequence, "text": p.text})
    for index, table in enumerate(document.tables, start=1):
        for row in table.rows:
            text = "\t".join(c.text.strip() for c in row.cells)
            if text.strip():
                sequence += 1
                paragraphs.append({"sequence": sequence, "text": text, "table": index})
    return paragraphs
