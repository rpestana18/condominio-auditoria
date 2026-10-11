"""Ensures the reader output follows the JSON Schema contract v2."""
import io
import json
from pathlib import Path

import jsonschema
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app

ROOT = Path(__file__).resolve().parents[2]
SCHEMA = json.loads((ROOT / "contracts/reader/v2/read-document.schema.json").read_text())
client = TestClient(app)


def send(name: str, content: bytes) -> dict:
    response = client.post("/v2/read", files={"file": (name, content)})
    assert response.status_code == 200, response.text
    body = response.json()
    jsonschema.validate(body, SCHEMA)
    return body


def test_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Rubrica"
    sheet["B1"] = 1234.56
    buffer = io.BytesIO()
    workbook.save(buffer)
    body = send("po.xlsx", buffer.getvalue())
    assert body["type"] == "xlsx"
    assert {"row": 1, "column": 2, "value": "1234.56", "type": "number"} in body["sheets"][0]["cells"]


def test_docx():
    document = Document()
    document.add_paragraph("Ata da assembleia")
    buffer = io.BytesIO()
    document.save(buffer)
    body = send("ata.docx", buffer.getvalue())
    assert body["type"] == "docx"
    assert body["paragraphs"][0]["text"] == "Ata da assembleia"


def test_sample_pdf_when_available():
    sample = ROOT / "data/golden/privado/fluxo-caixa-2026-09.pdf"
    if not sample.exists():
        import pytest
        pytest.skip("PDF de exemplo não está no repositório (dado real)")
    body = send(sample.name, sample.read_bytes())
    assert body["type"] == "pdf"
    assert len(body["pages"]) == 22


def test_unsupported_format():
    response = client.post("/v2/read", files={"file": ("x.txt", b"ola")})
    assert response.status_code == 415


def test_health():
    assert client.get("/health").json()["status"] == "ok"
